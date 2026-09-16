import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'C:\Users\phamm\Desktop\BDS_Hommy_Presentation_v2.pptx')
print(f'Total slides: {len(prs.slides)}')
print()

checks = {
    'Slide 1 GVHD': ('ThS. Phạm Đình Tài', 0),
    'Slide 1 ngành': ('Kỹ thuật Công nghệ Thông tin', 0),
    'Slide 3 bối cảnh': ('quy trình BĐS', 2),
    'Slide 4 mục tiêu': ('31 tài khoản', 3),
    'Slide 4 tech stack': ('Node.js v20', 3),
    'Slide 4 test count': ('40 Unit Test', 3),
    'Slide 6 Admin số liệu': ('31 TK', 5),
    'Slide 6 cuộc hẹn': ('22 cuộc hẹn', 5),
    'Slide 12 ML kết quả': ('89%', 11),
    'Slide 14 kết luận': ('kiểm thử thực tế', 13),
    'Slide 15 GVHD': ('ThS. Phạm Đình Tài', 14),
    'Slide 15 ngành': ('Kỹ thuật Công nghệ Thông tin', 14),
}

for check_name, (find_text, slide_idx) in checks.items():
    slide = prs.slides[slide_idx]
    found = False
    for shape in slide.shapes:
        if shape.has_text_frame and find_text.lower() in shape.text_frame.text.lower():
            found = True
            break
    status = 'OK' if found else 'MISSING'
    print(f'[{status}] {check_name}: "{find_text}"')
