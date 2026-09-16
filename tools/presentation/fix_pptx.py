"""
Bổ sung thêm:
1. Tên trường Đại học Nguyễn Tất Thành + Khoa CNTT vào Slide 1
2. Fix Slide 3 bối cảnh (dòng 2 và 3 bị thiếu)
3. Bổ sung tên trường vào Slide 15
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INPUT_FILE = r"C:\Users\phamm\Desktop\BDS_Hommy_Presentation_v2.pptx"
OUTPUT_FILE = r"C:\Users\phamm\Desktop\BDS_Hommy_Presentation_v2.pptx"

NAVY = RGBColor(0x00, 0x20, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xC0, 0x00)

prs = Presentation(INPUT_FILE)

def add_textbox(slide, left_in, top_in, width_in, height_in, text, 
                font_size=13, bold=False, color=None, italic=False, alignment=PP_ALIGN.LEFT):
    from pptx.util import Inches, Pt
    txBox = slide.shapes.add_textbox(
        Inches(left_in), Inches(top_in), 
        Inches(width_in), Inches(height_in)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = alignment
    run = para.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return txBox

# ============================================================
# SLIDE 1 — Thêm tên trường (nếu chưa có)
# ============================================================
slide1 = prs.slides[0]
print("Kiểm tra Slide 1 shapes:")

has_school = False
for shape in slide1.shapes:
    if shape.has_text_frame:
        txt = shape.text_frame.text
        print(f"  {shape.name}: {txt.strip()[:60]}")
        if "Nguyễn Tất Thành" in txt or "NTTU" in txt:
            has_school = True

if not has_school:
    print("-> Thêm tên trường vào Slide 1...")
    # Thêm textbox tên trường ở phía trên (dưới dòng tiêu đề đầu)
    # Vị trí phù hợp: khoảng row 2 (y~0.85in), căn giữa
    txbox_school = add_textbox(
        slide1, 
        left_in=0.5, top_in=0.88, 
        width_in=12.33, height_in=0.45,
        text="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH — KHOA KỸ THUẬT CÔNG NGHỆ THÔNG TIN",
        font_size=12,
        bold=False,
        color=NAVY,
        alignment=PP_ALIGN.CENTER
    )
    print("  -> Đã thêm tên trường Slide 1")
else:
    print("-> Đã có tên trường")

# ============================================================
# SLIDE 3 — Fix bối cảnh (các bullets bị thiếu do text quá dài)
# ============================================================
slide3 = prs.slides[2]
print("\nKiểm tra Slide 3:")
for shape in slide3.shapes:
    if shape.has_text_frame:
        txt = shape.text_frame.text.strip()
        if txt:
            print(f"  {shape.name}: {txt[:100]}")

# Tìm và cập nhật TextBox 8 (bullets bối cảnh 1) - chỉ còn 1 bullet
for shape in slide3.shapes:
    if shape.name == "TextBox 8" and shape.has_text_frame:
        tf = shape.text_frame
        current_count = len([p for p in tf.paragraphs if p.text.strip()])
        print(f"  TextBox 8 has {current_count} non-empty paragraphs")
        
        if current_count < 3:
            print("  -> Bổ sung bullet bị thiếu...")
            # Thêm các bullets còn thiếu
            bullets_to_add = [
                "  ►  Thiếu nền tảng số tích hợp đầy đủ: đăng tin, kiểm duyệt, đặt lịch, hợp đồng và thanh toán đang dùng các công cụ rời rạc",
                "  ►  Khó kiểm soát luồng tiền, hoa hồng và xác thực danh tính — tiềm ẩn rủi ro gian lận bất động sản"
            ]
            for bullet in bullets_to_add:
                para = tf.add_paragraph()
                run = para.add_run()
                run.text = bullet
                run.font.size = Pt(13)
                run.font.bold = False
                run.font.color.rgb = NAVY
            print("  -> Đã thêm 2 bullets bối cảnh")

# ============================================================
# SLIDE 15 — Thêm tên trường nếu chưa có
# ============================================================
slide15 = prs.slides[14]
print("\nKiểm tra Slide 15:")

has_school_15 = False
for shape in slide15.shapes:
    if shape.has_text_frame:
        txt = shape.text_frame.text
        print(f"  {shape.name}: {txt.strip()[:60]}")
        if "Nguyễn Tất Thành" in txt:
            has_school_15 = True

if not has_school_15:
    print("-> Thêm tên trường vào Slide 15...")
    txbox_school15 = add_textbox(
        slide15,
        left_in=0.5, top_in=6.15,
        width_in=12.33, height_in=0.40,
        text="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH — KHOA KỸ THUẬT CÔNG NGHỆ THÔNG TIN",
        font_size=11,
        bold=False,
        color=NAVY,
        alignment=PP_ALIGN.CENTER
    )
    print("  -> Đã thêm tên trường Slide 15")

# ============================================================
# LƯU FILE
# ============================================================
print("\nLưu file...")
prs.save(OUTPUT_FILE)
print(f"✅ Saved: {OUTPUT_FILE}")
